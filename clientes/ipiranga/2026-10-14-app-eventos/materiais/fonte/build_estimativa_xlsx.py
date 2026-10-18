#!/usr/bin/env python3
"""Modelo interno de estimativa: App Eventos Ipiranga (14/10/2026)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

INK = "13141D"; PINK = "F72585"; GREY = "6B6E7A"; BG = "F3F3F5"
thin = Side(style="thin", color="E3E4E8")
hdr_font = Font(name="Poppins", bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor=INK)
pink_fill = PatternFill("solid", fgColor=PINK)
bg_fill = PatternFill("solid", fgColor=BG)
base = Font(name="Poppins", size=10)
bold = Font(name="Poppins", size=10, bold=True)
title = Font(name="Poppins", size=14, bold=True, color=INK)
note = Font(name="Poppins", size=9, italic=True, color=GREY)

wb = Workbook()

# ---------------------------------------------------------------- Premissas
ws = wb.active; ws.title = "Premissas"
ws["A1"] = "Ipiranga · App Eventos · premissas de precificação"; ws["A1"].font = title
rows = [
    ("Valor hora (R$)", 200, "Mesma base do App de Execução (out/2026)"),
    ("Reserva técnica (% da execução)", 0.10, "Imprevistos de integração e revisão das lojas"),
    ("Direção (% da execução)", 0.30, "Bruno (produto/UX), Nóbrega (arquitetura), gestão, reuniões"),
    ("Fator total sobre execução", "=1+B3+B4", "Horas faturadas = execução × fator"),
    ("Horas produtivas por semana (time)", 60, "2 devs mobile/back em paralelo + design, média"),
]
ws.append([]); ws.append(["Parâmetro", "Valor", "Observação"])
for c in "ABC": ws[f"{c}3"].font = hdr_font; ws[f"{c}3"].fill = hdr_fill
# rows start at 4? No: A1 title, row 2 blank, row 3 header -> data rows 4..8. Adjust formulas.
for i, (k, v, o) in enumerate(rows):
    r = 4 + i
    ws.cell(r, 1, k).font = bold; ws.cell(r, 2, v).font = base; ws.cell(r, 3, o).font = note
ws["B5"].number_format = "0%"; ws["B6"].number_format = "0%"
ws["B7"] = "=1+B5+B6"; ws["B7"].number_format = "0.00"
ws.column_dimensions["A"].width = 38; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 70
ws["A10"] = "Como ler: cada aba de opção lista as horas de execução por frente. A reserva e a direção são calculadas sobre o total de execução; o preço é horas totais × valor hora."
ws["A10"].font = note; ws.merge_cells("A10:C10"); ws["A10"].alignment = Alignment(wrap_text=True); ws.row_dimensions[10].height = 32

HORA = "Premissas!$B$4"; RES = "Premissas!$B$5"; DIR = "Premissas!$B$6"; SEM = "Premissas!$B$8"

def option_sheet(name, heading, frentes, prazo_txt):
    s = wb.create_sheet(name)
    s["A1"] = heading; s["A1"].font = title
    s["A2"] = prazo_txt; s["A2"].font = note
    s.append([])
    s.append(["Frente", "Entregas", "Horas de execução", "R$ execução"])
    for c in "ABCD": s[f"{c}4"].font = hdr_font; s[f"{c}4"].fill = hdr_fill
    r0 = 5
    for i, (f, d, h) in enumerate(frentes):
        r = r0 + i
        s.cell(r, 1, f).font = bold; s.cell(r, 2, d).font = base
        s.cell(r, 3, h).font = base; s.cell(r, 4, f"=C{r}*{HORA}").number_format = '"R$" #,##0'
        s.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        for c in range(1, 5): s.cell(r, c).border = Border(bottom=thin); s.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
    rl = r0 + len(frentes) - 1
    r = rl + 1
    s.cell(r, 1, "Execução").font = bold; s.cell(r, 3, f"=SUM(C{r0}:C{rl})").font = bold
    s.cell(r, 4, f"=C{r}*{HORA}").font = bold
    s.cell(r + 1, 1, "Reserva técnica").font = base; s.cell(r + 1, 3, f"=C{r}*{RES}"); s.cell(r + 1, 4, f"=C{r+1}*{HORA}")
    s.cell(r + 2, 1, "Direção e gestão").font = base; s.cell(r + 2, 3, f"=C{r}*{DIR}"); s.cell(r + 2, 4, f"=C{r+2}*{HORA}")
    s.cell(r + 3, 1, "TOTAL").font = Font(name="Poppins", bold=True, color="FFFFFF", size=11)
    s.cell(r + 3, 3, f"=SUM(C{r}:C{r+2})").font = Font(name="Poppins", bold=True, color="FFFFFF", size=11)
    s.cell(r + 3, 4, f"=C{r+3}*{HORA}").font = Font(name="Poppins", bold=True, color="FFFFFF", size=11)
    for c in range(1, 5): s.cell(r + 3, c).fill = pink_fill
    s.cell(r + 4, 1, "Preço sugerido (arredondado)").font = bold
    s.cell(r + 4, 4, f"=ROUND(D{r+3}/1000,0)*1000").font = bold
    s.cell(r + 5, 1, "Prazo estimado (semanas)").font = bold
    s.cell(r + 5, 3, f"=ROUNDUP(C{r+3}/{SEM},0)").font = bold
    s.cell(r + 6, 1, "Pagamento sugerido").font = bold
    s.cell(r + 7, 1, "30% na assinatura").font = base; s.cell(r + 7, 4, f"=D{r+4}*0.3")
    s.cell(r + 8, 1, "30% na homologação (TestFlight / APK interno)").font = base; s.cell(r + 8, 4, f"=D{r+4}*0.3")
    s.cell(r + 9, 1, "40% na publicação nas lojas").font = base; s.cell(r + 9, 4, f"=D{r+4}*0.4")
    for rr in range(r, r + 10):
        for c in (3,):
            s.cell(rr, c).number_format = "0"
        s.cell(rr, 4).number_format = '"R$" #,##0'
    s.column_dimensions["A"].width = 34; s.column_dimensions["B"].width = 78
    s.column_dimensions["C"].width = 18; s.column_dimensions["D"].width = 16
    s.freeze_panes = "A5"
    return s

A = [
    ("Descoberta e arquitetura", "Mapeamento da API que o portal precisa expor (login, evento, agenda, mapa, inscrição em arena), definição da stack híbrida (React Native/Expo), fluxo de dados offline, plano de push", 24),
    ("UX/UI mobile", "Design system mobile a partir da identidade do portal; cerca de 12 telas: onboarding/login, home do evento, agenda por dia com favoritos, detalhe de atividade, mapa do evento com pontos, inscrição em arena, perfil/credencial, central de avisos; protótipo navegável", 60),
    ("Backend e API no portal", "Autenticação por token, endpoints de evento, agenda, mapa, inscrição em arena, cadastro de dispositivo; painel de envio de push no admin do portal (texto, público, agendamento)", 60),
    ("App mobile (iOS e Android)", "Setup do projeto, login integrado, home, agenda com favoritos e lembretes, mapa com zoom e pontos de interesse, inscrição em arena, perfil, central de avisos, cache offline de agenda e mapa, recebimento de push, deep link por e-mail/QR", 150),
    ("Push notification", "Integração FCM/APNs, registro de dispositivos, segmentação por evento e perfil (revendedor/funcionário), histórico de envios", 20),
    ("QA e homologação", "Testes em dispositivos iOS e Android, cenários sem conexão, matriz de versões de sistema, correções; rodada de homologação com a Ipiranga via TestFlight e APK", 36),
    ("Publicação nas lojas", "Contas de desenvolvedor da Ipiranga, assets de loja, política de privacidade, formulários de dados, submissão e resposta às revisões Apple e Google", 20),
]
B = [
    ("Descoberta e arquitetura", "Tudo da opção A, mais modelagem dos formulários dinâmicos de inscrição do portal (campos, regras, anexos, passagens e hotelaria quando houver), lista de eventos e perfis de acesso", 32),
    ("UX/UI mobile", "Cerca de 22 telas: além das telas da opção A, lista e filtro de eventos, detalhe de evento, fluxo completo de inscrição com formulários dinâmicos, meus eventos e histórico, documentos e materiais do evento, preferências de notificação", 100),
    ("Backend e API no portal", "API completa: catálogo de eventos, inscrições com formulários dinâmicos e validações, materiais e documentos, perfil e histórico, notificações por evento; painel de push por evento e por público", 110),
    ("App mobile (iOS e Android)", "Tudo da opção A para qualquer evento, renderizador de formulários dinâmicos, upload de documentos, meus eventos, materiais, preferências; compatível com a estrutura atual de eventos do portal sem retrabalho por evento", 260),
    ("Push notification", "Integração FCM/APNs, segmentação por evento, público e etapa da jornada (inscrito, confirmado, presente), agendamento e histórico", 30),
    ("QA e homologação", "Matriz maior de cenários (vários eventos, formulários distintos, offline), testes em dispositivos, rodadas de homologação com a Ipiranga", 56),
    ("Publicação nas lojas", "Igual à opção A", 20),
]
option_sheet("Opcao A - Convencao", "Opção A · App da Convenção (um evento, escopo reduzido)", A, "App dedicado ao evento principal, reaproveitável para a próxima convenção com troca de conteúdo pelo portal.")
option_sheet("Opcao B - Portal completo", "Opção B · App de todos os eventos do portal", B, "Um app permanente da Ipiranga Eventos, que acompanha qualquer evento criado no portal.")

# ---------------------------------------------------------------- PWA referência
s = wb.create_sheet("Referencia PWA")
s["A1"] = "Referência · evolução PWA do portal (caminho em teste)"; s["A1"].font = title
s["A2"] = "Não é a estimativa pedida; serve para comparar o custo de levar o portal até o limite do que o PWA permite."; s["A2"].font = note
s.append([]); s.append(["Frente", "Entregas", "Horas de execução", "R$ execução"])
for c in "ABCD": s[f"{c}4"].font = hdr_font; s[f"{c}4"].fill = hdr_fill
P = [
    ("Manifesto e instalação", "Ícone na tela inicial, splash, modo standalone, orientação de instalação no iOS e Android dentro do fluxo de inscrição", 12),
    ("Offline", "Service worker com cache de agenda, mapa e credencial do evento; estados sem conexão", 20),
    ("Interface mobile do evento", "Revisão da experiência no celular: home do evento, agenda, mapa com zoom, atalhos; login que cai direto no evento", 40),
    ("Salvar no calendário", "Função fixa no portal e na mensagem de sucesso (já em andamento, fora desta conta)", 0),
    ("Piloto com a estrutura da Convenção 2025", "Levantar o evento em homologação, rodada de uso com a Ipiranga, ajustes", 16),
]
for i, (f, d, h) in enumerate(P):
    r = 5 + i
    s.cell(r, 1, f).font = bold; s.cell(r, 2, d).font = base; s.cell(r, 3, h); s.cell(r, 4, f"=C{r}*{HORA}").number_format = '"R$" #,##0'
    s.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
r = 5 + len(P)
s.cell(r, 1, "Execução").font = bold; s.cell(r, 3, f"=SUM(C5:C{r-1})").font = bold; s.cell(r, 4, f"=C{r}*{HORA}")
s.cell(r+1, 1, "Com reserva e direção").font = bold; s.cell(r+1, 3, f"=C{r}*Premissas!$B$7").font = bold; s.cell(r+1, 4, f"=C{r+1}*{HORA}").font = bold
for rr in (r, r+1): s.cell(rr, 4).number_format = '"R$" #,##0'; s.cell(rr, 3).number_format = "0"
s.column_dimensions["A"].width = 34; s.column_dimensions["B"].width = 78; s.column_dimensions["C"].width = 18; s.column_dimensions["D"].width = 16

# ---------------------------------------------------------------- Recorrentes
s = wb.create_sheet("Recorrentes")
s["A1"] = "Custos recorrentes de um app nativo (fora do investimento de construção)"; s["A1"].font = title
s.append([]); s.append(["Item", "Quem paga", "Valor", "Periodicidade", "Observação"])
for c in "ABCDE": s[f"{c}3"].font = hdr_font; s[f"{c}3"].fill = hdr_fill
R = [
    ("Apple Developer Program", "Ipiranga", "US$ 99", "anual", "Conta da empresa, exige D-U-N-S; prazo de aprovação da conta de 1 a 3 semanas"),
    ("Google Play Console", "Ipiranga", "US$ 25", "única", "Conta da empresa; verificação de identidade"),
    ("Push (Firebase Cloud Messaging)", "Ipiranga", "sem custo", "mensal", "Volume de eventos cabe no plano gratuito"),
    ("Sustentação do app", "Fattoria", "=10*Premissas!$B$4", "mensal, mínimo", "10 h/mês: atualizações de sistema operacional, bibliotecas, revisões de loja, pequenos ajustes. Picos de evento fora da base"),
    ("Sustentação do app", "Fattoria", "=16*Premissas!$B$4", "mensal, confortável", "16 h/mês cobre também evolução leve entre eventos"),
    ("Hospedagem e API do portal", "Ipiranga", "incluída", "mensal", "Já coberta pela operação atual do portal; a API acrescenta carga marginal"),
]
for i, row in enumerate(R):
    r = 4 + i
    for c, v in enumerate(row, 1):
        s.cell(r, c, v).font = base; s.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
    s.cell(r, 3).number_format = '"R$" #,##0'
for c, w in zip("ABCDE", (32, 12, 14, 20, 70)): s.column_dimensions[c].width = w

# ---------------------------------------------------------------- Resumo
s = wb.create_sheet("Resumo", 0)
s["A1"] = "Ipiranga · App Eventos · resumo da estimativa (14/10/2026)"; s["A1"].font = title
s.append([]); s.append(["Caminho", "Horas totais", "Investimento", "Preço sugerido", "Prazo (semanas)", "O que resolve"])
for c in "ABCDEF": s[f"{c}3"].font = hdr_font; s[f"{c}3"].fill = hdr_fill
# find total rows: option sheets: data starts row 5, len(A)=7 -> exec row 12, total row 15, preço 16, prazo 17
def refs(sheet, n):
    exec_r = 5 + n; tot = exec_r + 3
    return f"='{sheet}'!C{tot}", f"='{sheet}'!D{tot}", f"='{sheet}'!D{tot+1}", f"='{sheet}'!C{tot+2}"
for i, (lab, sheet, n, txt) in enumerate([
    ("Referência · PWA do portal", None, None, "Praticidade: ícone na tela, agenda e mapa offline, acesso direto ao evento. Sem push."),
    ("Opção A · App da Convenção", "Opcao A - Convencao", len(A), "Tudo do PWA mais push notification e experiência nativa, para um evento por vez."),
    ("Opção B · App Ipiranga Eventos", "Opcao B - Portal completo", len(B), "App permanente para todos os eventos, inscrição completa pelo celular, push por evento e público."),
]):
    r = 4 + i
    s.cell(r, 1, lab).font = bold; s.cell(r, 6, txt).font = base
    if sheet:
        h, v, p, w = refs(sheet, n)
        s.cell(r, 2, h); s.cell(r, 3, v); s.cell(r, 4, p); s.cell(r, 5, w)
    else:
        pr = 5 + len(P) + 1
        s.cell(r, 2, f"='Referencia PWA'!C{pr}"); s.cell(r, 3, f"='Referencia PWA'!D{pr}"); s.cell(r, 4, f"=ROUND(C{r}/1000,0)*1000"); s.cell(r, 5, 3)
    for c in range(1, 7): s.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top"); s.cell(r, c).border = Border(bottom=thin)
    s.cell(r, 2).number_format = "0"; s.cell(r, 3).number_format = '"R$" #,##0'; s.cell(r, 4).number_format = '"R$" #,##0'
for c, w in zip("ABCDEF", (32, 13, 15, 15, 15, 70)): s.column_dimensions[c].width = w
s["A9"] = "Prazos contam a partir do kick-off e não incluem o tempo de abertura das contas de desenvolvedor da Ipiranga (1 a 3 semanas, em paralelo) nem a revisão das lojas na primeira publicação (2 a 7 dias)."
s["A9"].font = note; s.merge_cells("A9:F9"); s["A9"].alignment = Alignment(wrap_text=True); s.row_dimensions[9].height = 30

import sys
out = sys.argv[1]
wb.save(out); print("ok", out)
