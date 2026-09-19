# -*- coding: utf-8 -*-
"""v2: sem notas. Desenha o visual do roadmap macro (cinco fases + passos futuros)."""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from content_v2 import SLIDES

PURPLE = RGBColor(0x6E, 0x5B, 0xD5)
LILAC = RGBColor(0xE4, 0xE0, 0xF7)
DARK = RGBColor(0x33, 0x33, 0x33)
GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation('stage_v2.pptx')
assert len(prs.slides) == len(SLIDES), (len(prs.slides), len(SLIDES))

def style_runs(tf, size, bold=False, color=DARK, font='Poppins', align=None):
    for p in tf.paragraphs:
        if align is not None:
            p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size); r.font.bold = bold; r.font.name = font
            r.font.color.rgb = color

def box(slide, x, y, w, h, text, fill, color, size, bold=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    shp.adjustments[0] = 0.16
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.fill.background(); shp.shadow.inherit = False
    tf = shp.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(50000); tf.margin_top = tf.margin_bottom = Emu(30000)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.text = text
    style_runs(tf, size, bold=bold, color=color, align=PP_ALIGN.CENTER)
    return shp

def label(slide, x, y, w, h, text, size=8, color=GREY, font='Poppins Medium', align=PP_ALIGN.CENTER, bold=False):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(30000); tf.margin_top = tf.margin_bottom = Emu(0)
    tf.text = text
    style_runs(tf, size, bold=bold, color=color, font=font, align=align)
    return tb

for slide, spec in zip(prs.slides, SLIDES):
    if spec.get('extra') != 'phases':
        continue
    x0, total_w = 379826, 8303700
    # linha 1: cinco fases
    steps = spec['phases']; n = len(steps); gap = 70000
    w = (total_w - gap * (n - 1)) // n
    y1 = 1250000
    for i, (name, desc) in enumerate(steps):
        x = x0 + i * (w + gap)
        box(slide, x, y1, w, 640000, name, PURPLE, WHITE, 10)
        label(slide, x, y1 + 680000, w, 520000, desc)
        if i < n - 1:
            label(slide, x + w - 15000, y1 + 190000, gap + 30000, 300000, '›', size=12, bold=True)
    # linha 2: passos futuros
    label(slide, x0, 2480000, total_w, 260000, 'PASSOS FUTUROS, DEPOIS DA PRIMEIRA COMPRA VALIDADA', size=8, color=GREY, font='Poppins', align=PP_ALIGN.LEFT, bold=True)
    fut = spec['future']; m = len(fut)
    wf = (total_w - gap * (m - 1)) // m
    y2 = 2760000
    for i, (name, _) in enumerate(fut):
        x = x0 + i * (wf + gap)
        box(slide, x, y2, wf, 520000, name, LILAC, PURPLE, 9)

prs.save('ICONIC_Roadmap_Estrategia_v2.pptx')
print('saved ICONIC_Roadmap_Estrategia_v2.pptx with', len(prs.slides), 'slides')
